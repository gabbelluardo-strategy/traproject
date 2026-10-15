import copy, sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

SRC, OUT = sys.argv[1], sys.argv[2]
prs = Presentation(SRC)

NAVY = RGBColor(0x24, 0x37, 0x82)
ORANGE = RGBColor(0xE9, 0x4E, 0x24)
TEXT = RGBColor(0x27, 0x2B, 0x35)
GREY = RGBColor(0x5E, 0x63, 0x70)
BORDER = RGBColor(0xDA, 0xDF, 0xE8)
TINT = RGBColor(0xEE, 0xF0, 0xFA)
TINT2 = RGBColor(0xE1, 0xE5, 0xF6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PEACH = RGBColor(0xFC, 0xE6, 0xDE)
BOLD = "Encode Sans Expanded"
LIGHT = "Encode Sans ExpandedLight"

tpl = list(prs.slides)
cover, agenda, blank, divider, cards_tpl, takeout = tpl
SLDNUM = next(s for s in cards_tpl.placeholders if s.placeholder_format.idx == 17)._element
LAYOUT_BAR = cards_tpl.slide_layout          # "2 columns with a photo": header bar
LAYOUT_DIAG = divider.slide_layout           # "Text with a graph": diagonal


def fmt_run(run, size, color=TEXT, bold=False, italic=False, font=None):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = color
    f.name = font or (BOLD if bold else LIGHT)


def text(slide, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT, name=None):
    """paras: list of (text, size, color, bold, italic, space_after) or list of runs."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    first = True
    for para in paras:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        runs = para["runs"] if isinstance(para, dict) else [para]
        space = para.get("after", 0) if isinstance(para, dict) else (para[5] if len(para) > 5 else 0)
        bullet = para.get("bullet", False) if isinstance(para, dict) else False
        for r in runs:
            t, size, color, bold, italic = r[:5]
            run = p.add_run()
            run.text = t
            fmt_run(run, size, color, bold, italic)
        p.space_after = Pt(space)
        if bullet:
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(Emu(Inches(0.18))))
            pPr.set("indent", str(-Emu(Inches(0.18))))
            bu = etree_sub(pPr, "a:buChar")
            bu.set("char", "•")
    return tb


def etree_sub(parent, tag):
    from lxml import etree
    for old in parent.findall(qn("a:buNone")):
        parent.remove(old)
    el = etree.SubElement(parent, qn(tag))
    return el


def box(slide, x, y, w, h, fill=WHITE, line=BORDER, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05, name=None):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(0.75)
    s.shadow.inherit = False
    if name:
        s.name = name
    s.text_frame.text = ""
    return s


def label_box(slide, x, y, w, h, t, fill=NAVY, color=WHITE, size=14, bold=False, shape=MSO_SHAPE.RECTANGLE, align=PP_ALIGN.CENTER, name=None):
    s = box(slide, x, y, w, h, fill=fill, line=None, shape=shape, radius=0.3, name=name)
    tf = s.text_frame
    tf.margin_left = tf.margin_right = Inches(0.08)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = t
    fmt_run(r, size, color, bold)
    return s


def new_slide(layout, header=None, header_idx=19, keep=()):
    s = prs.slides.add_slide(layout)
    for ph in list(s.placeholders):
        idx = ph.placeholder_format.idx
        if idx == header_idx and header is not None:
            ph.text_frame.text = header
            for r in ph.text_frame.paragraphs[0].runs:
                r.font.size = Pt(16)
        elif idx not in keep:
            ph._element.getparent().remove(ph._element)
    s.shapes._spTree.append(copy.deepcopy(SLDNUM))
    return s


def notes(slide, t):
    slide.notes_slide.notes_text_frame.text = t


def title(slide, t, y=1.12, size=22, w=12.0):
    return text(slide, 0.67, y, w, 0.6, [(t, size, NAVY, True, False)], name="Title")


def footnote(slide, t):
    text(slide, 0.67, 6.98, 11.2, 0.3, [(t, 9, GREY, False, True)], name="Footnote")


def vertical_label(slide, t="ESG Narrative"):
    tb = text(slide, -0.95, 3.9, 2.6, 0.3, [(t, 12, NAVY, True, False)], align=PP_ALIGN.CENTER, name="Side label")
    tb.rotation = 270


# ---------------------------------------------------------------- 1 cover
for sh in cover.shapes:
    if sh.has_text_frame and sh.text_frame.text.startswith("ESG"):
        paras = sh.text_frame.paragraphs
        last = paras[-1]
        newp = copy.deepcopy(last._p)
        last._p.addnext(newp)
        runs = newp.findall(qn("a:r"))
        for r in runs[1:]:
            newp.remove(r)
        runs[0].find(qn("a:t")).text = "Global sustainability narrative: corporate storytelling"
        rPr = runs[0].find(qn("a:rPr"))
        rPr.set("sz", "1800")
        rPr.set("b", "0")
        rPr.set("cap", "none")
        pPr = newp.find(qn("a:pPr"))
        sb = pPr.find(qn("a:spcBef"))
        sb.clear(); from lxml import etree as _e; _e.SubElement(sb, qn("a:spcPts")).set("val", "1400")
        sh.top = sh.top - Inches(0.35)
        sh.height = sh.height + Inches(0.35)
    elif sh.has_text_frame and sh.text_frame.text.strip() == "2026":
        sh.text_frame.paragraphs[0].runs[0].text = "October 2026"
notes(cover, "Third deep-dive on the global sustainability narrative, focused on overall corporate storytelling: organising idea, narrative arc, hub architecture, stakeholder layers, link to business strategy, consistency over time. Reference case named by the client team: Volkswagen Group's sustainability hub.")

# ---------------------------------------------------------------- 2 agenda
for sh in list(agenda.shapes):
    if sh.has_text_frame and sh.text_frame.text.startswith("Da trad"):
        sh._element.getparent().remove(sh._element)
    elif sh.has_text_frame and sh.text_frame.text == "AMBITI DI RICERCA":
        sh.text_frame.paragraphs[0].runs[0].text = "RESEARCH AREAS"
    elif sh.has_text_frame and sh.text_frame.text.strip() == "Lorem":
        tf = sh.text_frame
        p0 = tf.paragraphs[0]
        p0.runs[0].text = "Global sustainability narrative: corporate storytelling"
        for hl in p0._p.iter(qn("a:highlight")):
            hl.getparent().remove(hl)
        sh.width = Inches(9.5)
        subs = ["Key insight", "Reference case: Volkswagen Group", "Best practices: automotive and beyond",
                "Benchmark overview and expert view", "Stellantis today", "Design principles and take-outs"]
        for t in subs:
            p = tf.add_paragraph()
            p.level = 1
            r = p.add_run()
            r.text = t
            fmt_run(r, 14, TEXT, False, font=LIGHT)
notes(agenda, "One research area in this deck: the global sustainability narrative, read through overall corporate storytelling. The sub-items are the sections of the deck.")

# ---------------------------------------------------------------- 3 divider
for sh in divider.shapes:
    if sh.is_placeholder and sh.placeholder_format.idx == 13:
        sh.text_frame.paragraphs[0].runs[0].text = "GLOBAL SUSTAINABILITY NARRATIVE"
    elif sh.is_placeholder and sh.placeholder_format.idx == 21:
        sh.text_frame.paragraphs[0].runs[0].text = "1"
    elif sh.name == "Text 35":
        ps = sh.text_frame.paragraphs
        ps[0].runs[0].text = "KEY QUESTION"
        ps[2].runs[0].text = ("How do top players link ESG and business strategy through one recognisable "
                              "story that stays coherent over time, across markets and stakeholders?")
        sh.width = Inches(6.4)
notes(divider, "Brief from the ESG and HR Communication teams: make sustainability communication simpler, more visual, more engaging and connected to business strategy. This section answers it through overall storytelling, with Volkswagen Group's hub as the reference.")

# ---------------------------------------------------------------- 4 key insight
s4 = new_slide(LAYOUT_BAR, "Key insight")
text(s4, 0.67, 1.2, 11.6, 1.3, [("The best sustainability storytellers do not have the smartest slogan. "
                                 "They have one organising idea, repeated identically wherever the company speaks.",
                                 24, NAVY, True, False)], name="Statement")
cards4 = [("01", "One grid, everywhere", "Hub, KPIs, report, governance and local sites all use the same dimensions."),
          ("02", "Impact areas, not dates", "Toyota, VW and Schneider survived target changes. Mercedes' Ambition 2039 and Stellantis' 2038 did not."),
          ("03", "Prove and admit", "One number for every promise, and the courage to say what is still missing.")]
for i, (n, h, b) in enumerate(cards4):
    x = 0.67 + i * 4.07
    box(s4, x, 3.0, 3.85, 2.75, name=f"Card {n}")
    text(s4, x + 0.3, 3.3, 3.25, 2.9, [
        {"runs": [(n, 28, ORANGE, True, False)], "after": 10},
        {"runs": [(h, 18, NAVY, True, False)], "after": 8},
        {"runs": [(b, 15, TEXT, False, False)]}], name=f"Card {n} text")
notes(s4, "In 2026 the consistency of a sustainability narrative is tested above all when a target changes. Almost every carmaker revised a target in 2024-2026: frameworks built on impact areas held, frameworks built on an end date suffered.")

# ---------------------------------------------------------------- 5 VW reference
s5 = new_slide(LAYOUT_BAR, "Reference case · Volkswagen Group")
title(s5, "regenerate+: one grid beyond the website", w=7.7)
label_box(s5, 0.67, 1.9, 7.6, 0.85, "", fill=NAVY, name="regenerate+ box")
text(s5, 0.85, 1.97, 7.3, 0.75, [
    {"runs": [("regenerate+", 18, WHITE, True, False)], "after": 2},
    {"runs": [("“A mobility provider with positive added value for nature and society”", 11, WHITE, False, True)]}],
    align=PP_ALIGN.CENTER, name="regenerate+ text")
dims = [("Nature", "3 Top-KPIs: decarbonisation, circularity, biodiversity"), ("Our People", "3 Top-KPIs"),
        ("Society", "3 Top-KPIs"), ("Business", "3 Top-KPIs: the most distinctive dimension")]
for i, (d, k) in enumerate(dims):
    x = 0.67 + i * 1.94
    box(s5, x, 2.95, 1.78, 1.45, name=f"Dimension {d}")
    text(s5, x + 0.14, 3.08, 1.5, 1.25, [{"runs": [(d, 13, NAVY, True, False)], "after": 4},
                                          {"runs": [(k, 10, TEXT, False, False)]}], name=f"Dimension {d} text")
text(s5, 0.67, 4.6, 7.6, 0.3, [("THE SAME GRID ORGANISES", 11, ORANGE, True, False)], name="Grid label")
uses = ["The web hub and strategy page", "The regenerate+ Progress Report", "The external council: 4 Practice Groups", "The Decarbonization Index in Board pay"]
for i, u in enumerate(uses):
    label_box(s5, 0.67 + i * 1.94, 4.95, 1.78, 1.0, u, fill=TINT, color=NAVY, size=10, bold=True, shape=MSO_SHAPE.ROUNDED_RECTANGLE, name=f"Use {i+1}")
box(s5, 8.6, 1.2, 4.1, 5.6, name="Assessment card")
label_box(s5, 8.6, 1.2, 4.1, 0.41, "Volkswagen Group", size=13, name="Logo box")
text(s5, 8.85, 1.85, 3.6, 4.7, [
    {"runs": [("Why it works", 10, ORANGE, True, False)], "after": 2},
    {"runs": [("A proprietary name, four dimensions with three Top-KPIs each, and one grid for site, report, external council and Board pay. "
               "“Elevate sustainability” sits inside the group strategy.", 11, TEXT, False, False)], "after": 10},
    {"runs": [("Why it matters for Stellantis", 10, ORANGE, True, False)], "after": 2},
    {"runs": [("Adopt the skeleton, then add what VW lacks: its own voice, an explicit link to FaSTLAne 2030 and a People dimension owned by HR (VW already has “Our People”).", 11, TEXT, False, False)], "after": 10},
    {"runs": [("What stays generic", 10, ORANGE, True, False)], "after": 2},
    {"runs": [("Textbook ESG names, a framework only about two years old, no verified visual signature.", 11, TEXT, False, False)]}],
    name="Assessment text")
text(s5, 10.0, 6.45, 2.5, 0.25, [("#OneGridEverywhere", 10, NAVY, True, True)], align=PP_ALIGN.RIGHT, name="Hashtag")
footnote(s5, "Sources: VW Sustainability hub, VW Annual Report 2024-2025, EMobility+, MarketScreener. Hub page not opened directly: visual judgement provisional.")
notes(s5, "VW's strength is architecture, not aesthetics. Each dimension has three Top-KPIs; since October 2024 the external Sustainability Council is split into four Practice Groups, one per dimension; next to the ESRS statement VW publishes a narrative report with the same name. An executive: “It has to make business sense, it's not philanthropic.” Progress figures circulating (Decarbonization Index 46.3 t CO2e per vehicle in 2025 vs 48.5 target) come from secondary sources. Visual to capture: top of https://www.volkswagen-group.com/en/sustainability-15772 with regenerate+ and the four dimensions; Top-KPI table in the ESRS 2025 PDF.")


# ---------------------------------------------------------------- case slides
def case_slide(header, cases, note):
    s = new_slide(LAYOUT_BAR, header)
    vertical_label(s)
    for i, c in enumerate(cases):
        x = 0.66 + i * 4.07
        box(s, x, 1.62, 3.9, 5.25, name=f"Card {i+1}")
        label_box(s, x, 1.13, 3.08, 0.41, c["logo"], size=14, name=f"Logo {i+1}")
        text(s, x + 0.22, 1.85, 3.46, 4.55, [
            {"runs": [("Title", 10, ORANGE, True, False)], "after": 2},
            {"runs": [(c["title"], 12, NAVY, True, False)], "after": 12},
            {"runs": [("Why it works", 10, ORANGE, True, False)], "after": 2},
            {"runs": [(c["why"], 11, TEXT, False, False)], "after": 12},
            {"runs": [("Why it matters for Stellantis", 10, ORANGE, True, False)], "after": 2},
            {"runs": [(c["matters"], 11, TEXT, False, False)]}], name=f"Card {i+1} text")
        text(s, x + 1.2, 6.5, 2.48, 0.25, [(c["tag"], 10, NAVY, True, True)], align=PP_ALIGN.RIGHT, name=f"Hashtag {i+1}")
    notes(s, note)
    return s


s6 = case_slide("Best practices · Automotive", [
    dict(logo="Mercedes-Benz", title="The deepest hub: from promise to single model",
         why="One path, five layers: overview, six strategic areas, Ambition 2039, a 360° Environmental Check per model, investor metrics and a Climate Transition Action Plan. Plant target −50% reached in 2022, then raised to −80%.",
         matters="The same layered path, from one screen to the investor PDF, plus a public Climate Transition Plan linked from the 2050 page, as EmpCo will require.",
         tag="#LayeredHub"),
    dict(logo="Toyota", title="Environmental Challenge 2050: eleven years, same frame",
         why="Six numbered challenges since 2015: impact areas, not deadlines. Interim dates change, the story does not. The same icon set travels across UK, US, Lebanese and dealer sites.",
         matters="Make the dimensions numbered, iconic and dateless, with dates kept in the KPIs. Proof that a multi-energy plan like FaSTLAne can carry a lasting story.",
         tag="#ImpactAreasNotDates"),
    dict(logo="Renault Group", title="The Future is NEUTRAL: sustainability in the brand menu",
         why="Circularity sits under “brands & entities” next to Renault, Dacia and Alpine: a business with its own logo and targets (over €2.3bn revenue, over 10% margin by 2030, self-reported). Born in 2022, kept in futuREady.",
         matters="Move SUSTAINera out of the “Vehicles” page and into the group's entities, with revenue and margin KPIs: proof that sustainability pays for every powertrain.",
         tag="#CircularityAsBusiness")],
    "Mercedes: hub, six strategic areas (Environment & Climate, Resources & Circularity, Human Rights, Traffic Safety, Digital Trust, People), Ambition 2039 since 2019, 360° Environmental Check per model, investor ESG page with CTAP 2025. Watch-out: six areas are hard to remember and Ambition 2039 is built on a date, already weakened in 2024. Visuals: group.mercedes-benz.com/sustainability/six-strategic-areas.html and the Ambition 2039 page.\n"
    "Toyota: New Vehicle Zero CO2, Life Cycle Zero CO2, Plant Zero CO2, Water, Recycling-based society, Society in harmony with nature. Watch-out: 2026 global hub not verified; NGO criticism of the multi-pathway strategy to be checked. Visual: toyota.co.uk Environmental Challenge 2050 icon set.\n"
    "Renault: today SUSTAINera is a line on the net zero “Vehicles” page with “up to 80%/50%” claims that EmpCo puts at risk; replace with specific LCA values. Visual: renaultgroup.com brands & entities menu.")

s7 = case_slide("Best practices · Beyond automotive", [
    dict(logo="Microsoft", title="Promise, proof, people, depth",
         why="Four 2030 commitments, each with one big proof number and its own sub-page, then the CSO's voice, stories and a reports hub. It owns the bad news in the first person: emissions up about 25% in 2025, solutions “not scaling fast enough”.",
         matters="The model for the first two layers of the hub: 3-4 commitments, one number and one page each. And the register for 2050: name the setback before the press does.",
         tag="#PromiseProofPeople"),
    dict(logo="Schneider Electric", title="A score out of 10 that survives a change of plan",
         why="SSI 2021-25 closed at 8.86/10; Impact 2030 restarts the scale (3.69 at H1 2026 vs a 4.20 target). Published quarterly with results, with pillars written as business missions, one of them for human potential.",
         matters="A “Stellantis Impact Score” that restarts with FaSTLAne turns the new plan into a chapter, not a retreat, and is reported with quarterly results.",
         tag="#OneScore"),
    dict(logo="Patagonia", title="“Work in Progress”: candour as a brand trait",
         why="A purpose, an act that proves it (“Earth is now our only shareholder”) and a first impact report with missed targets on the cover: footprint up 2%. Ranked first by GlobeScan/ERM experts in 2025.",
         matters="One lesson only: name the step back, explain why, reaffirm the path. The opposite of Unilever's quiet 2024 target cuts. The tone itself is not transferable.",
         tag="#RadicalCandour")],
    "Microsoft: the only hub we could read in full (October 2026). Proof numbers are self-reported (100% renewable electricity match, 14.2m m3 water replenished, 92% server reuse and recycling). Watch-out: criticism over a gas deal with Chevron; a good hub does not close the say-do gap. Visual: microsoft.com/en-us/corporate-responsibility/sustainability hero and progress cards.\n"
    "Schneider: 862 MtCO2 avoided for customers vs an 800 Mt ambition (self-reported). Watch-out: weight in bonuses not verified; composite indices need public sub-indicators.\n"
    "Patagonia: FY25 footprint 182,646 tCO2e (+2%); 100% preferred-materials goal missed (about 84-86%). Chief impact officer: “a transparency tool, not a victory lap”. Watch-out: candour works only when backed by a concrete act.")

# ---------------------------------------------------------------- 8 overview table
s8 = new_slide(LAYOUT_BAR, "Benchmark overview")
title(s8, "Depth beats slogans; impact areas beat dates")
rows = [("Case", "Organising idea", "Years", "Link to business", "Verdict"),
        ("Volkswagen", "regenerate+, 4 dimensions × 3 Top-KPIs", "~2", "Index in Board pay", "Top: architecture"),
        ("Mercedes-Benz", "6 areas + Ambition 2039", "7", "Investor ESG event", "Top: depth"),
        ("Toyota", "Challenge 2050, 6 challenges", "11", "Multi-pathway Beyond Zero", "Top: longevity"),
        ("Renault", "The Future is NEUTRAL", "4", "€2.3bn revenue, >10% margin", "Top: business link"),
        ("Microsoft", "Four 2030 commitments", "Several", "AI and cloud, $1bn fund", "Top: architecture"),
        ("Schneider Electric", "Impact score out of 10", "Several cycles", "Quarterly, with results", "Top: across plans"),
        ("Patagonia", "Purpose + “Work in Progress”", "Since 2022", "The company is the act", "Top: credibility"),
        ("BYD", "“Cool the Earth by 1°C”", "~2017", "EV product", "Solid: memorable"),
        ("Stellantis", "None", "n/a", "FaSTLAne with no story", "Generic, inconsistent")]
widths = [2.0, 3.6, 1.4, 2.9, 2.1]
tbl = s8.shapes.add_table(len(rows), 5, Inches(0.67), Inches(1.85), Inches(sum(widths)), Inches(0.42 * len(rows))).table
for j, w in enumerate(widths):
    tbl.columns[j].width = Inches(w)
for i, row in enumerate(rows):
    tbl.rows[i].height = Inches(0.42)
    for j, val in enumerate(row):
        c = tbl.cell(i, j)
        c.fill.solid()
        if i == 0:
            c.fill.fore_color.rgb = NAVY
        elif row[0] == "Stellantis":
            c.fill.fore_color.rgb = PEACH
        else:
            c.fill.fore_color.rgb = TINT if i % 2 == 0 else WHITE
        c.margin_left = c.margin_right = Inches(0.1)
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf = c.text_frame
        tf.text = ""
        r = tf.paragraphs[0].add_run()
        r.text = val
        fmt_run(r, 11, WHITE if i == 0 else (NAVY if j == 0 else TEXT), bold=(i == 0 or j == 0))
footnote(s8, "Years and verdicts from secondary sources; launch years of VW, Microsoft, Schneider and BYD not verified. The Stellantis verdict is our inference.")
notes(s8, "Frameworks built on impact areas (Toyota, VW, Schneider) survive target changes better than those built on an end date (Ambition 2039, net zero 2038). The 2024-2026 credibility hits landed on dates: Mercedes and Volvo in 2024, Polestar 0 from 2030 to 2035, Stellantis from 2038 to 2050. A slogan alone (BYD) is the most memorable but the thinnest in depth.")

# ---------------------------------------------------------------- 9 rankings matrix
s9 = new_slide(LAYOUT_BAR, "What rankings and experts say")
title(s9, "Substance first, distinctiveness right after")
mx, my, mw, mh = 0.67, 2.2, 6.4, 4.2
text(s9, mx, 1.85, mw, 0.3, [("↑ Distinctiveness: a credible, authentic story", 11, NAVY, True, False)], name="Y axis")
quads = [(0, 0, TINT, "Slogan without depth"), (1, 0, TINT2, "Substance with a voice"),
         (0, 1, TINT, "Neither"), (1, 1, TINT, "Substance without a voice")]
for qx, qy, fill, lab in quads:
    x, y = mx + qx * mw / 2, my + qy * mh / 2
    box(s9, x, y, mw / 2 - 0.04, mh / 2 - 0.04, fill=fill, line=None, shape=MSO_SHAPE.RECTANGLE, name=f"Quadrant {lab}")
    ly = y + 0.1 if qy == 0 else y + mh / 2 - 0.4
    text(s9, x + 0.12, ly, mw / 2 - 0.3, 0.3, [(lab, 10, GREY, False, True)], name=f"Quadrant label {lab}")
pts = [("BYD", 0.35, 0.75, NAVY), ("Patagonia", 4.65, 0.45, NAVY), ("Microsoft", 3.55, 0.95, NAVY), ("Toyota", 4.75, 1.45, NAVY),
       ("Volkswagen", 4.35, 2.55, NAVY), ("Mercedes-Benz", 3.5, 3.1, NAVY), ("Stellantis", 0.35, 2.85, ORANGE)]
for n, px, py, col in pts:
    label_box(s9, mx + px, my + py, 0.13 * len(n) + 0.35, 0.34, n, fill=col, size=10, bold=True, shape=MSO_SHAPE.ROUNDED_RECTANGLE, name=f"Point {n}")
text(s9, mx, my + mh + 0.08, mw, 0.3, [("Substance → a concrete, structured, transparent journey", 11, NAVY, True, False)], name="X axis")
facts = [("Lundquist", "The “CSRD paradox”: more mandatory reporting means more need to stand out. In 2025 it calls for a Sustainability Value Proposition."),
         ("Webranking 2025-26", "Poste Italiane first. Social sustainability lags environmental across the board: natural space for HR."),
         ("GlobeScan/ERM 2025", "Patagonia, IKEA, Unilever and Natura &Co lead; no carmaker among the names cited.")]
for i, (h, b) in enumerate(facts):
    y = 1.85 + i * 1.6
    box(s9, 7.45, y, 5.2, 1.45, name=f"Fact card {i+1}")
    text(s9, 7.7, y + 0.18, 4.75, 1.15, [{"runs": [(h, 13, NAVY, True, False)], "after": 4},
                                         {"runs": [(b, 11, TEXT, False, False)]}], name=f"Fact text {i+1}")
footnote(s9, "Axes from Lundquist; positioning of the cases is our qualitative reading, not a Lundquist score.")
notes(s9, "Lundquist rates sustainability communication on two axes: substance (“the concreteness of a structured, transparent and integrated journey”) and distinctiveness (“the ability to tell a credible and authentic story”). Webranking: about 50 criteria, European average 47.1/100; Iberdrola sixth with 87.3 and an almost full sustainability section. WBCSD Reporting Matters 2025: a chapter on design and storytelling; 48% of reports CSRD-aligned. Bain: CEO rhetoric is shifting from moral value to business value.")

# ---------------------------------------------------------------- 10 context
s10 = new_slide(LAYOUT_BAR, "Context 2024-2026")
title(s10, "Between silence and empty promises, EmpCo asks for a public plan")
cols = [("Greenhushing", ["85% of 75 large companies kept or expanded their programmes; only 16% said so (reported by HBR, second-hand).", "South Pole 2024: found in 9 sectors out of 14."], False),
        ("Say-do gap", ["Marsh: be candid on status, targets and plan; silence attracts more criticism.", "Bain: CEOs are moving from moral value to business value."], False),
        ("EmpCo, from 27 Sept 2026", ["“Net zero by 2050” needs a public plan with targets, resources and governance.", "The plan must be reviewed by an independent expert, with public results."], True)]
for i, (h, items, dark) in enumerate(cols):
    x = 0.67 + i * 4.07
    box(s10, x, 1.9, 3.85, 3.6, fill=NAVY if dark else WHITE, line=None if dark else BORDER, name=f"Column {i+1}")
    paras = [{"runs": [(h, 16, WHITE if dark else NAVY, True, False)], "after": 12}]
    for it in items:
        paras.append({"runs": [(it, 14, WHITE if dark else TEXT, False, False)], "after": 12, "bullet": True})
    text(s10, x + 0.3, 2.2, 3.25, 3.6, paras, name=f"Column {i+1} text")
text(s10, 0.67, 5.75, 12.0, 0.6, [("To verify: fines (consultancy sources), transposition in 7 of 27 Member States by March 2026, applicability to corporate hubs. The CSRD Omnibus (December 2025) keeps Stellantis in scope.", 11, GREY, False, False)], name="Caveat")
notes(s10, "With EmpCo, corporate storytelling becomes a legal matter. The most relevant rule is the one on future-performance claims: a detailed, realistic public implementation plan with measurable targets, resources, governance and monitoring, regularly reviewed by an independent expert. The plan can sit on the corporate site, linked to the claim. The prudent reading is that the corporate hub is in scope.")

# ---------------------------------------------------------------- 11 Stellantis facts
s11 = new_slide(LAYOUT_BAR, "Stellantis today · The facts")
title(s11, "Plenty of content, no story")
stats = [("2050", "Net zero moved from 2038, yet pages with both years are still online, Brazil included"),
         ("−43%", "Scope 1-2 in 2025 vs 2021 (−26% Scope 1-3): proof points that do not make a story today"),
         ("€60bn", "FaSTLAne 2030 “to accelerate growth and profit”: no decarbonisation target in the announcements"),
         ("32nd → 96th", "Reputation Manager Italy, H1 to Q4 2025; Lead the Charge from 5th to 9th")]
for i, (n, b) in enumerate(stats):
    x = 0.67 + i * 3.05
    dark = i == 3
    box(s11, x, 1.95, 2.85, 2.6, fill=NAVY if dark else WHITE, line=None if dark else BORDER, name=f"Stat {i+1}")
    text(s11, x + 0.25, 2.2, 2.4, 2.5, [{"runs": [(n, 28, ORANGE if dark else NAVY, True, False)], "after": 10},
                                        {"runs": [(b, 13, WHITE if dark else TEXT, False, False)]}], name=f"Stat {i+1} text")
text(s11, 0.67, 4.9, 12.0, 1.0, [{"runs": [("The hub is organised by stakeholder (customers, employees, host countries, suppliers, planet) and by a three-area net zero strategy. The press filled the narrative gap with one word: ", 14, TEXT, False, False),
                                             ("“delay”.", 14, ORANGE, True, False)]}], name="Hub note")
footnote(s11, "Sources: Stellantis, MarketScreener, ClubAlfa, Milano Finanza, Primaonline, Lead the Charge. stellantis.com could not be opened directly.")
notes(s11, "Navigation, visuals and localisation of stellantis.com were not assessed directly. Pages still indexed with 2038: “CSR Vision”, “An Ambitious Decarbonization Strategy” and the Brazilian mirror; not verified whether updated after May 2026. The 2025 Climate Policy Report explains the reset with “market dynamics and government policy”. Milano Finanza links it to write-downs of over €22bn.")

# ---------------------------------------------------------------- 12 Stellantis gaps
s12 = new_slide(LAYOUT_BAR, "Stellantis today · The diagnosis")
title(s12, "Six storytelling gaps (our inference)")
gaps = [("No organising idea", "Stakeholder lists and technical pillars are not a story."),
        ("Broken consistency", "Two net zero years online; local sites keep the old content."),
        ("FaSTLAne disconnected", "It does not explain why a multi-energy plan can be sustainable."),
        ("No “why”", "The reset is there; a first-person explanation is not."),
        ("HR story underused", "Motor Citizens (6,000 colleagues, 23,000 hours in 2024) and Shares to Win stay on the margins."),
        ("Reputation not helped", "Falling indices say sustainability is not an asset of the story today.")]
for i, (h, b) in enumerate(gaps):
    x = 0.67 + (i % 3) * 4.07
    y = 1.85 + (i // 3) * 2.1
    box(s12, x, y, 3.85, 1.9, name=f"Gap {i+1}")
    text(s12, x + 0.25, y + 0.2, 3.35, 1.55, [{"runs": [(f"{i+1}  ", 18, ORANGE, True, False), (h, 14, NAVY, True, False)], "after": 6},
                                               {"runs": [(b, 13, TEXT, False, False)]}], name=f"Gap {i+1} text")
text(s12, 0.67, 6.15, 12.0, 0.6, [("EmpCo adds urgency: “net zero 2050” and SUSTAINera's “up to 80%/50%” claims need a public plan and specific proof.", 13, NAVY, True, False)], name="EmpCo note")
notes(s12, "Our own diagnosis, built on secondary sources: the Stellantis hub was not opened. The social story is exactly where, according to Comprend, all companies lag: a natural space for HR Communication.")

# ---------------------------------------------------------------- 13 principles
s13 = new_slide(LAYOUT_BAR, "Design principles")
title(s13, "Four principles for touchpoint-based communication")
pr = [("01", "Proof at the moment of choice", "Navigation, configurator, parts counter, energy contract (Google Maps, Back Market, Renault V2G)."),
      ("02", "One programme, one object, one link", "Every commitment becomes an artefact with a QR code to verifiable data (Kia, Hyundai Motorstudio, L'Oréal)."),
      ("03", "Euros, kilometres and kilos, not adjectives", "It holds under EmpCo, works for every FaSTLAne powertrain and does not depend on the net zero date."),
      ("04", "Long platform, local chapters, people", "A multi-year frame renewed by each brand and market, told by technicians and employees.")]
for i, (n, h, b) in enumerate(pr):
    x = 0.67 + (i % 2) * 6.1
    y = 1.95 + (i // 2) * 2.15
    box(s13, x, y, 5.9, 1.95, name=f"Principle {n}")
    label_box(s13, x + 0.3, y + 0.3, 0.75, 0.75, n, fill=NAVY, size=18, bold=True, shape=MSO_SHAPE.OVAL, name=f"Principle {n} badge")
    text(s13, x + 1.3, y + 0.3, 4.35, 1.75, [{"runs": [(h, 16, NAVY, True, False)], "after": 6},
                                             {"runs": [(b, 14, TEXT, False, False)]}], name=f"Principle {n} text")
footnote(s13, "From the companion benchmark on sustainability touchpoints.")
notes(s13, "These four principles come from the companion touchpoint benchmark and turn it into design criteria. The fourth answers the request for consistency over time: IONIQ Forest and LVMH LIFE 360 are multi-year group frames that each brand and market renews.")

# ---------------------------------------------------------------- 14 take-outs
for sh in takeout.shapes:
    if sh.is_placeholder and sh.placeholder_format.idx == 13:
        sh.text_frame.paragraphs[0].runs[0].text = "TAKE-OUTS"
    elif sh.is_placeholder and sh.placeholder_format.idx == 21:
        sh.text_frame.paragraphs[0].runs[0].text = "1"
    elif sh.name == "Text 35":
        tf = sh.text_frame
        p0 = tf.paragraphs[0]
        p0.runs[0].text = "The 2050 reset is not an obstacle to the story: it is the chance to build one that does not depend on a date."
        p0.runs[0].font.size = Pt(20)
        p0.runs[0].font.color.rgb = NAVY
        p0.runs[0].font.bold = True
        for extra in tf.paragraphs[1:]:
            extra._p.getparent().remove(extra._p)
        p0.space_after = Pt(14)
        for t in ["One organising idea and one grid, everywhere the group speaks.",
                  "Impact areas in the framework; dates in the KPIs and the transition plan.",
                  "Prove every claim and say what is still missing.",
                  "A framework where ESG guarantees the proof and HR provides the people."]:
            p = tf.add_paragraph()
            r = p.add_run()
            r.text = t
            fmt_run(r, 15, TEXT, False)
            p.space_after = Pt(8)
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(Emu(Inches(0.25))))
            pPr.set("indent", str(-Emu(Inches(0.25))))
            etree_sub(pPr, "a:buChar").set("char", "•")
        sh.width = Inches(7.0)
        sh.height = Inches(3.6)
notes(takeout, "With EmpCo, corporate storytelling becomes a legal matter; with the Omnibus the report becomes a strategic tool again; Webranking shows the social side is the weakest everywhere. Together they open a space that Stellantis' ESG and HR teams can occupy jointly. Before presenting: capture the visuals by hand (only Microsoft's hub could be opened), and check the ESG content of the FaSTLAne Investor Day deck.")

# ---------------------------------------------------------------- order & cleanup
order = [cover, agenda, divider, s4, s5, s6, s7, s8, s9, s10, s11, s12, s13, takeout]
sldIdLst = prs.slides._sldIdLst
id_by_slide = {}
for sldId in list(sldIdLst):
    id_by_slide[prs.part.related_part(sldId.rId)] = sldId
for drop in (blank, cards_tpl):
    sid = id_by_slide[drop.part]
    prs.part.drop_rel(sid.rId)
    sldIdLst.remove(sid)
for s in order:
    sid = id_by_slide[s.part]
    sldIdLst.remove(sid)
    sldIdLst.append(sid)
prs.save(OUT)
print("saved", OUT)
