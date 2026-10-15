import copy, sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

SRC, OUT = sys.argv[1], sys.argv[2]
prs = Presentation(SRC)

NAVY = RGBColor(0x24, 0x37, 0x82)
ORANGE = RGBColor(0xE9, 0x4E, 0x24)
TEXT = RGBColor(0x27, 0x2B, 0x35)
GREY = RGBColor(0x5E, 0x63, 0x70)
BORDER = RGBColor(0xDA, 0xDF, 0xE8)
TINT = RGBColor(0xEE, 0xF0, 0xFA)
TINT2 = RGBColor(0xE1, 0xE5, 0xF6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PEACH = RGBColor(0xFC, 0xE6, 0xDE)
BOLD = "Encode Sans Expanded"
LIGHT = "Encode Sans ExpandedLight"

tpl = list(prs.slides)
cover, agenda, blank, divider, cards_tpl, takeout = tpl
SLDNUM = next(s for s in cards_tpl.placeholders if s.placeholder_format.idx == 17)._element
LAYOUT_BAR = cards_tpl.slide_layout          # "2 columns with a photo": header bar
LAYOUT_DIAG = divider.slide_layout           # "Text with a graph": diagonal


def fmt_run(run, size, color=TEXT, bold=False, italic=False, font=None):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = color
    f.name = font or (BOLD if bold else LIGHT)


def text(slide, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT, name=None):
    """paras: list of (text, size, color, bold, italic, space_after) or list of runs."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    first = True
    for para in paras:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        runs = para["runs"] if isinstance(para, dict) else [para]
        space = para.get("after", 0) if isinstance(para, dict) else (para[5] if len(para) > 5 else 0)
        bullet = para.get("bullet", False) if isinstance(para, dict) else False
        for r in runs:
            t, size, color, bold, italic = r[:5]
            run = p.add_run()
            run.text = t
            fmt_run(run, size, color, bold, italic)
        p.space_after = Pt(space)
        if bullet:
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(Emu(Inches(0.18))))
            pPr.set("indent", str(-Emu(Inches(0.18))))
            bu = etree_sub(pPr, "a:buChar")
            bu.set("char", "•")
    return tb


def etree_sub(parent, tag):
    from lxml import etree
    for old in parent.findall(qn("a:buNone")):
        parent.remove(old)
    el = etree.SubElement(parent, qn(tag))
    return el


def box(slide, x, y, w, h, fill=WHITE, line=BORDER, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05, name=None):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(0.75)
    s.shadow.inherit = False
    if name:
        s.name = name
    s.text_frame.text = ""
    return s


def label_box(slide, x, y, w, h, t, fill=NAVY, color=WHITE, size=14, bold=False, shape=MSO_SHAPE.RECTANGLE, align=PP_ALIGN.CENTER, name=None):
    s = box(slide, x, y, w, h, fill=fill, line=None, shape=shape, radius=0.3, name=name)
    tf = s.text_frame
    tf.margin_left = tf.margin_right = Inches(0.08)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = t
    fmt_run(r, size, color, bold)
    return s


def new_slide(layout, header=None, header_idx=19, keep=()):
    s = prs.slides.add_slide(layout)
    for ph in list(s.placeholders):
        idx = ph.placeholder_format.idx
        if idx == header_idx and header is not None:
            ph.text_frame.text = header
            for r in ph.text_frame.paragraphs[0].runs:
                r.font.size = Pt(16)
        elif idx not in keep:
            ph._element.getparent().remove(ph._element)
    s.shapes._spTree.append(copy.deepcopy(SLDNUM))
    return s


def notes(slide, t):
    slide.notes_slide.notes_text_frame.text = t


def title(slide, t, y=1.12, size=22, w=12.0):
    return text(slide, 0.67, y, w, 0.6, [(t, size, NAVY, True, False)], name="Title")


def footnote(slide, t):
    text(slide, 0.67, 6.98, 11.2, 0.3, [(t, 9, GREY, False, True)], name="Footnote")


def vertical_label(slide, t="ESG Narrative"):
    tb = text(slide, -0.95, 3.9, 2.6, 0.3, [(t, 12, NAVY, True, False)], align=PP_ALIGN.CENTER, name="Side label")
    tb.rotation = 270


